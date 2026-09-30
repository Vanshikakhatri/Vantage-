"""
ReportGenerator — writes evaluation results to a clean, single-sheet-per-view
Excel workbook. Run `python -m src.eval.run_eval` (see that file) after
querying your test set to produce data/reports/vantage_eval_report.xlsx.

Sheets:
  1. Query Results   — one row per test question, all metrics + confidence
  2. Confusion Matrix — TP/FP/TN/FN + precision/recall/F1/accuracy (formulas)
  3. Summary          — headline averages or the full run (formulas)
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.eval.metrics import EvalRecord, confusion_matrix

_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
_BODY_FONT = Font(name="Arial", size=10)
_GOOD_FILL = PatternFill("solid", fgColor="D1FAE5")
_WARN_FILL = PatternFill("solid", fgColor="FEF3C7")
_BAD_FILL = PatternFill("solid", fgColor="FEE2E2")


def _style_header(ws, row: int, n_cols: int) -> None:
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws, n_cols: int, width: int = 18) -> None:
    for col in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = width


class ReportGenerator:
    def generate(self, records: list[EvalRecord], output_path: Path,
                 compliance_frameworks: list[str] | None = None) -> Path:
        wb = Workbook()
        self._write_query_results(wb.active, records)
        self._write_confusion_matrix(wb.create_sheet("Confusion Matrix"), records)
        self._write_summary(wb.create_sheet("Summary"), records, compliance_frameworks or [])

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        return output_path

    def _write_query_results(self, ws, records: list[EvalRecord]) -> None:
        ws.title = "Query Results"
        headers = ["Query ID", "Question", "Confidence /10", "Faithfulness",
                   "Hallucination Rate", "Answer Relevancy", "Context Precision",
                   "Groundedness (cosine)", "Sources Retrieved", "Predicted Grounded",
                   "Expected Grounded", "Latency (ms)"]
        ws.append(headers)
        _style_header(ws, 1, len(headers))

        for r in records:
            ws.append([
                r.query_id, r.question, r.confidence_score, r.faithfulness,
                r.hallucination_rate, r.answer_relevancy, r.context_precision,
                r.groundedness_cosine, r.num_sources_retrieved, r.predicted_grounded,
                r.expected_grounded if r.expected_grounded is not None else "—", r.latency_ms,
            ])
            row = ws.max_row
            for cell in ws[row]:
                cell.font = _BODY_FONT
            hall_cell = ws.cell(row=row, column=5)
            hall_cell.fill = (_GOOD_FILL if r.hallucination_rate <= 0.15 else
                               _WARN_FILL if r.hallucination_rate <= 0.35 else _BAD_FILL)

        for col in (4, 5, 6, 7, 8):
            for row in range(2, ws.max_row + 1):
                ws.cell(row=row, column=col).number_format = "0.00"

        _autosize(ws, len(headers))
        ws.column_dimensions["B"].width = 45
        ws.freeze_panes = "A2"

    def _write_confusion_matrix(self, ws, records: list[EvalRecord]) -> None:
        headers = ["Metric", "Value"]
        ws.append(["Confusion Matrix — Predicted Grounded vs Expected Grounded"])
        ws["A1"].font = Font(name="Arial", bold=True, size=13)
        ws.append([])
        ws.append(headers)
        _style_header(ws, 3, 2)

        n_rows = len(records) + 1  # +1 for header row on "Query Results" sheet
        # Formulas reference the Query Results sheet directly so the matrix
        # recalculates automatically if you paste in a new eval run.
        qr = "'Query Results'"
        rows = [
            ("True Positives (TP)",
             f'=SUMPRODUCT((({qr}!J2:J{n_rows}=1)*({qr}!K2:K{n_rows}=1)))'),
            ("False Positives (FP)",
             f'=SUMPRODUCT((({qr}!J2:J{n_rows}=1)*({qr}!K2:K{n_rows}=0)))'),
            ("True Negatives (TN)",
             f'=SUMPRODUCT((({qr}!J2:J{n_rows}=0)*({qr}!K2:K{n_rows}=0)))'),
            ("False Negatives (FN)",
             f'=SUMPRODUCT((({qr}!J2:J{n_rows}=0)*({qr}!K2:K{n_rows}=1)))'),
        ]
        start_row = ws.max_row + 1
        for label, formula in rows:
            ws.append([label, formula])
            ws.cell(row=ws.max_row, column=1).font = _BODY_FONT
            ws.cell(row=ws.max_row, column=2).font = _BODY_FONT

        tp_row, fp_row, tn_row, fn_row = start_row, start_row + 1, start_row + 2, start_row + 3
        ws.append([])
        derived = [
            ("Precision = TP/(TP+FP)", f"=IFERROR(B{tp_row}/(B{tp_row}+B{fp_row}),0)"),
            ("Recall = TP/(TP+FN)", f"=IFERROR(B{tp_row}/(B{tp_row}+B{fn_row}),0)"),
            ("F1 = 2PR/(P+R)", None),   # filled below, needs the precision/recall rows
            ("Accuracy = (TP+TN)/Total", f"=IFERROR((B{tp_row}+B{tn_row})/(B{tp_row}+B{fp_row}+B{tn_row}+B{fn_row}),0)"),
        ]
        prec_row = ws.max_row + 2
        for label, formula in derived:
            ws.append([label, formula])
        f1_row = prec_row + 2
        ws.cell(row=f1_row, column=2).value = (
            f"=IFERROR(2*B{prec_row}*B{prec_row+1}/(B{prec_row}+B{prec_row+1}),0)"
        )
        for row in range(4, ws.max_row + 1):
            for col in (1, 2):
                cell = ws.cell(row=row, column=col)
                if cell.value is not None:
                    cell.font = _BODY_FONT
        _autosize(ws, 2, 34)

    def _write_summary(self, ws, records: list[EvalRecord], frameworks: list[str]) -> None:
        ws.append(["Vantage — Evaluation Summary"])
        ws["A1"].font = Font(name="Arial", bold=True, size=14)
        ws.append([f"Compliance frameworks checked: {', '.join(frameworks) or 'none specified'}"])
        ws.append([f"Total queries evaluated: {len(records)}"])
        ws.append([])

        n = len(records) + 1
        qr = "'Query Results'"
        ws.append(["Metric", "Average", "Interpretation"])
        _style_header(ws, ws.max_row, 3)
        metrics_rows = [
            ("Avg. Confidence Score (/10)", f"=AVERAGE({qr}!C2:C{n})", "≥7.5 High · 5–7.5 Medium · <5 Low"),
            ("Avg. Faithfulness (0-1)", f"=AVERAGE({qr}!D2:D{n})", "Higher = fewer unsupported claims"),
            ("Avg. Hallucination Rate", f"=AVERAGE({qr}!E2:E{n})", "Lower is better; target < 0.15"),
            ("Avg. Answer Relevancy (0-1)", f"=AVERAGE({qr}!F2:F{n})", "Does the answer address the question"),
            ("Avg. Context Precision (0-1)", f"=AVERAGE({qr}!G2:G{n})", "Retrieval quality — noise vs signal"),
            ("Avg. Groundedness Cosine (0-1)", f"=AVERAGE({qr}!H2:H{n})", "Deterministic cross-check on faithfulness"),
            ("Avg. Latency (ms)", f"=AVERAGE({qr}!L2:L{n})", "End-to-end pipeline latency"),
        ]
        for label, formula, note in metrics_rows:
            ws.append([label, formula, note])
            ws.cell(row=ws.max_row, column=2).number_format = "0.00"
            for c in ws[ws.max_row]:
                c.font = _BODY_FONT
        _autosize(ws, 3, 34)
        ws.column_dimensions["C"].width = 46
